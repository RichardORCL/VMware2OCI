from __future__ import annotations

import hashlib
import io
import json
import tempfile
import unittest
import zipfile
from copy import deepcopy
from datetime import date, timedelta
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

from services.ocvs_sdd_adapter import normalize_sdd_configuration
from services.ocvs_sdd_delivery import mark_delivered, store_final_artifacts
from services.ocvs_sdd_handover import (
    action_overdue,
    build_handover_package,
    build_migration_planning_payload,
    confirm_handover,
    create_action,
    default_checklist,
    find_acceptance,
    invalidate_for_new_revision,
    normalize_acceptance_handover,
    readiness_decision,
    recalculate_readiness,
    record_acceptance_decision,
    set_readiness_owners,
    start_acceptance,
    summary,
    update_action,
    update_checklist_item,
)
from tests.test_ocvs_sdd_delivery import DOCX, PDF, finalized_configuration


class OCVSSDDHandoverTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name) / "artifacts"
        self.config, self.snapshot = finalized_configuration()

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def delivered(self) -> tuple[dict, str]:
        config = deepcopy(self.config)
        config["delivery_governance"] = store_final_artifacts(self.root, config, self.snapshot, DOCX, PDF, "Alex")
        artifact_id = config["delivery_governance"]["versions"][0]["artifact_id"]
        config, errors = mark_delivered(
            config, self.root, artifact_id, "Alex", "2026-09-17",
            {"name": "Casey", "company": "Northstar", "email": "casey@example.org"},
        )
        self.assertEqual(errors, [])
        return config, artifact_id

    def started(self) -> tuple[dict, str, str]:
        config, artifact_id = self.delivered()
        config, errors = start_acceptance(config, self.root, artifact_id, "Alex")
        self.assertEqual(errors, [])
        acceptance_id = config["acceptance_handover"]["active_acceptance_id"]
        return config, artifact_id, acceptance_id

    def decided(self, decision: str = "accepted", conditions: str = "") -> tuple[dict, str, str]:
        config, artifact_id, acceptance_id = self.started()
        fields = {
            "decision": decision, "customer_representative": "Casey Customer", "customer_role": "Sponsor",
            "customer_company": "Northstar", "decision_date": "2026-09-18", "customer_email": "casey@example.org",
            "comments": "Approved" if decision != "rejected" else "Design change required",
            "conditions": conditions, "evidence_reference": "SIGN-001", "recorded_by": "Alex",
            "recording_date": "2026-09-18",
        }
        config, errors = record_acceptance_decision(config, self.root, acceptance_id, fields)
        self.assertEqual(errors, [])
        return config, artifact_id, acceptance_id

    def ready(self, decision: str = "accepted") -> tuple[dict, str, str]:
        conditions = "Confirm change window" if decision == "accepted_with_conditions" else ""
        config, artifact_id, acceptance_id = self.decided(decision, conditions)
        governance = normalize_acceptance_handover(config["acceptance_handover"])
        record = next((item for item in governance["records"] if item["id"] == acceptance_id), None)
        assert record is not None
        for item in record["checklist"]:
            item["status"] = "complete"
        for action in record["actions"]:
            action["owner"] = "Project Manager"
            action["priority"] = "medium"
        record["implementation_owner"] = "Implementation Lead"
        record["migration_planning_owner"] = "Migration Lead"
        config["acceptance_handover"] = governance
        config = recalculate_readiness(config, self.root, acceptance_id)
        return config, artifact_id, acceptance_id

    def handed_over(self) -> tuple[dict, str, str]:
        config, artifact_id, acceptance_id = self.ready()
        config, errors = confirm_handover(config, self.root, acceptance_id, {
            "handed_over_by": "Alex", "received_by": "Morgan", "receiving_company": "Delivery Team",
            "handover_date": "2026-09-19", "implementation_owner": "Implementation Lead",
            "migration_planning_owner": "Migration Lead", "planned_start_date": "2026-10-01", "comments": "Ready",
        })
        self.assertEqual(errors, [])
        return config, artifact_id, acceptance_id

    def test_01_old_assessments_load_safe_defaults(self):
        self.assertEqual(normalize_sdd_configuration({})["acceptance_handover"]["records"], [])

    def test_02_only_delivered_versions_enter_acceptance(self):
        config = deepcopy(self.config)
        config["delivery_governance"] = store_final_artifacts(self.root, config, self.snapshot, DOCX, PDF, "Alex")
        artifact_id = config["delivery_governance"]["versions"][0]["artifact_id"]
        _config, errors = start_acceptance(config, self.root, artifact_id, "Alex")
        self.assertTrue(errors)

    def test_03_superseded_version_cannot_be_accepted(self):
        config, artifact_id = self.delivered()
        config["delivery_governance"]["versions"][0]["status"] = "superseded"
        _config, errors = start_acceptance(config, self.root, artifact_id, "Alex")
        self.assertTrue(errors)

    def test_04_accepted_requires_mandatory_fields(self):
        config, _artifact_id, acceptance_id = self.started()
        _config, errors = record_acceptance_decision(config, self.root, acceptance_id, {"decision": "accepted"})
        self.assertGreaterEqual(len(errors), 4)

    def test_05_accepted_with_conditions_requires_condition(self):
        config, _artifact_id, acceptance_id = self.started()
        fields = {"decision":"accepted_with_conditions","customer_representative":"C","customer_company":"N","decision_date":"2026-09-18","recorded_by":"A","recording_date":"2026-09-18"}
        _config, errors = record_acceptance_decision(config, self.root, acceptance_id, fields)
        self.assertTrue(any("condition" in error.lower() for error in errors))

    def test_06_rejected_requires_reason(self):
        config, _artifact_id, acceptance_id = self.started()
        fields = {"decision":"rejected","customer_representative":"C","customer_company":"N","decision_date":"2026-09-18","recorded_by":"A","recording_date":"2026-09-18"}
        _config, errors = record_acceptance_decision(config, self.root, acceptance_id, fields)
        self.assertTrue(any("reason" in error.lower() for error in errors))

    def test_07_invalid_email_is_rejected(self):
        config, _artifact_id, acceptance_id = self.started()
        fields = {"decision":"accepted","customer_representative":"C","customer_company":"N","decision_date":"2026-09-18","recorded_by":"A","recording_date":"2026-09-18","customer_email":"invalid"}
        _config, errors = record_acceptance_decision(config, self.root, acceptance_id, fields)
        self.assertTrue(any("e-mail" in error for error in errors))

    def test_08_duplicate_acceptance_is_prevented(self):
        config, artifact_id, _acceptance_id = self.started()
        _config, errors = start_acceptance(config, self.root, artifact_id, "Alex")
        self.assertTrue(any("already exists" in error for error in errors))

    def test_09_acceptance_references_immutable_version(self):
        config, artifact_id, acceptance_id = self.started()
        record = find_acceptance(config["acceptance_handover"], acceptance_id=acceptance_id)
        self.assertEqual((record["artifact_id"], record["sdd_version"]), (artifact_id, "1.0"))

    def test_10_delivered_artifacts_remain_unchanged(self):
        config, artifact_id, _acceptance_id = self.decided()
        artifact = config["delivery_governance"]["versions"][0]
        self.assertEqual(hashlib.sha256((self.root/artifact_id/artifact["docx_filename"]).read_bytes()).hexdigest(), artifact["docx_checksum"])

    def test_11_integrity_failure_blocks_acceptance(self):
        config, artifact_id = self.delivered(); artifact = config["delivery_governance"]["versions"][0]
        (self.root/artifact_id/artifact["pdf_filename"]).write_bytes(b"tampered")
        _config, errors = start_acceptance(config, self.root, artifact_id, "Alex")
        self.assertTrue(any("integrity" in error.lower() for error in errors))

    def test_12_checklist_status_persists(self):
        config, _artifact, acceptance_id = self.started(); item = config["acceptance_handover"]["records"][0]["checklist"][1]
        config, errors = update_checklist_item(config, self.root, acceptance_id, item["id"], {"status":"in_progress"}, "Alex")
        self.assertEqual(errors, []); self.assertEqual(config["acceptance_handover"]["records"][0]["checklist"][1]["status"], "in_progress")

    def test_13_checklist_owner_and_date_persist(self):
        config, _artifact, acceptance_id = self.started(); item = config["acceptance_handover"]["records"][0]["checklist"][1]
        config, _ = update_checklist_item(config, self.root, acceptance_id, item["id"], {"status":"complete","owner":"Sam","target_date":"2026-10-01"}, "Alex")
        updated = config["acceptance_handover"]["records"][0]["checklist"][1]
        self.assertEqual((updated["owner"], updated["target_date"]), ("Sam", "2026-10-01"))

    def test_14_blocked_checklist_prevents_readiness(self):
        config, _artifact, acceptance_id = self.decided(); record = config["acceptance_handover"]["records"][0]
        self.assertEqual(readiness_decision(config, self.root, record)["status"], "Not ready")

    def test_15_not_applicable_does_not_block_readiness(self):
        config, _artifact, acceptance_id = self.decided(); governance = normalize_acceptance_handover(config["acceptance_handover"]); record = governance["records"][0]
        for item in record["checklist"]: item["status"] = "not_applicable"
        record["implementation_owner"] = "A"; record["migration_planning_owner"] = "B"; config["acceptance_handover"] = governance
        self.assertEqual(readiness_decision(config, self.root, record)["status"], "Ready for implementation")

    def test_16_conditions_create_actions(self):
        config, _artifact, acceptance_id = self.decided("accepted_with_conditions", "Condition one\nCondition two")
        record = find_acceptance(config["acceptance_handover"], acceptance_id=acceptance_id)
        self.assertEqual([a["source"] for a in record["actions"]], ["customer_condition", "customer_condition"])

    def test_17_blocking_action_prevents_handover(self):
        config, _artifact, acceptance_id = self.ready(); config, _ = create_action(config, self.root, acceptance_id, {"description":"Blocker","priority":"blocking"}, "Alex")
        record = find_acceptance(config["acceptance_handover"], acceptance_id=acceptance_id)
        self.assertEqual(record["readiness"]["status"], "Not ready")

    def test_18_closed_actions_remain_in_history(self):
        config, _artifact, acceptance_id = self.ready(); config, _ = create_action(config, self.root, acceptance_id, {"description":"Task"}, "Alex")
        action = config["acceptance_handover"]["records"][0]["actions"][0]
        config, _ = update_action(config, self.root, acceptance_id, action["id"], {"status":"closed"}, "Alex")
        self.assertEqual(config["acceptance_handover"]["records"][0]["actions"][0]["status"], "closed")

    def test_19_overdue_actions_are_identified(self):
        self.assertTrue(action_overdue({"status":"open","due_date":(date.today()-timedelta(days=1)).isoformat()}))

    def test_20_ready_with_conditions_is_calculated(self):
        config, _artifact, acceptance_id = self.ready("accepted_with_conditions")
        self.assertEqual(find_acceptance(config["acceptance_handover"], acceptance_id=acceptance_id)["readiness"]["status"], "Ready with conditions")

    def test_21_ready_for_implementation_is_calculated(self):
        config, _artifact, acceptance_id = self.ready()
        self.assertEqual(find_acceptance(config["acceptance_handover"], acceptance_id=acceptance_id)["readiness"]["status"], "Ready for implementation")

    def test_22_handover_requires_ownership_fields(self):
        config, _artifact, acceptance_id = self.ready(); _config, errors = confirm_handover(config, self.root, acceptance_id, {})
        self.assertTrue(any("owner" in error.lower() for error in errors))

    def test_23_handover_creates_audit_event(self):
        config, _artifact, _acceptance_id = self.handed_over()
        self.assertIn("implementation_handover_confirmed", [e["event"] for e in config["acceptance_handover"]["audit_events"]])

    def test_24_accepted_version_becomes_baseline(self):
        config, artifact_id, acceptance_id = self.handed_over(); record = find_acceptance(config["acceptance_handover"], acceptance_id=acceptance_id)
        self.assertEqual(record["handover"]["baseline_artifact_id"], artifact_id)

    def test_25_handover_zip_contains_nine_files(self):
        config, _artifact, acceptance_id = self.handed_over(); package, filename, _manifest = build_handover_package(self.root, config, self.snapshot, acceptance_id)
        self.assertTrue(filename.endswith("_Implementation_Handover.zip"))
        with zipfile.ZipFile(io.BytesIO(package)) as archive: self.assertEqual(len(archive.namelist()), 9)

    def test_26_packaged_payload_files_have_checksums(self):
        config, _artifact, acceptance_id = self.handed_over(); package, _filename, manifest = build_handover_package(self.root, config, self.snapshot, acceptance_id)
        with zipfile.ZipFile(io.BytesIO(package)) as archive:
            payload_names = set(archive.namelist()) - {"handover-manifest.json"}
        self.assertEqual(payload_names, set(manifest["artifacts"]))

    def test_27_manifests_have_no_absolute_paths(self):
        config, _artifact, acceptance_id = self.handed_over(); package, _filename, manifest = build_handover_package(self.root, config, self.snapshot, acceptance_id)
        self.assertNotIn(str(self.root), json.dumps(manifest)); self.assertNotIn(str(self.root), package.decode("latin1", errors="ignore"))

    def test_28_new_revision_invalidates_acceptance(self):
        config, _artifact, _acceptance_id = self.ready(); config = invalidate_for_new_revision(config, "Alex", "1.1")
        self.assertTrue(config["acceptance_handover"]["records"][0]["baseline_superseded"])

    def test_29_previous_acceptance_details_are_retained(self):
        config, _artifact, _acceptance_id = self.ready(); before = deepcopy(config["acceptance_handover"]["records"][0]["acceptance"])
        config = invalidate_for_new_revision(config, "Alex", "1.1")
        self.assertEqual(config["acceptance_handover"]["records"][0]["acceptance"], before)

    def test_30_migration_payload_uses_selected_scope(self):
        config, _artifact, acceptance_id = self.handed_over(); self.snapshot["selected_vm_names"] = ["vm-1", "vm-2"]
        payload, errors = build_migration_planning_payload(config, self.snapshot, acceptance_id)
        self.assertEqual(errors, []); self.assertEqual(payload["selected_vm_names"], ["vm-1", "vm-2"])

    def test_31_workspace_is_scoped_to_ocvs_in_routes(self):
        source = Path("app.py").read_text(encoding="utf-8")
        self.assertIn('Customer Acceptance & Handover is available only for Move to OCVS.', source)

    def test_32_mobile_layout_has_responsive_rules(self):
        css = Path("static/css/sdd-acceptance.css").read_text(encoding="utf-8")
        self.assertIn("@media", css); self.assertIn("grid-template-columns:1fr", css)

    def test_33_generated_pdf_and_xlsx_open(self):
        config, _artifact, acceptance_id = self.handed_over(); package, _filename, _manifest = build_handover_package(self.root, config, self.snapshot, acceptance_id)
        with zipfile.ZipFile(io.BytesIO(package)) as archive:
            pdf = archive.read("implementation-handover-summary.pdf")
            self.assertEqual(len(PdfReader(io.BytesIO(pdf)).pages), 1)
            for name in ("implementation-readiness-checklist.xlsx", "open-action-register.xlsx"):
                workbook = load_workbook(io.BytesIO(archive.read(name)), read_only=True)
                self.assertEqual(workbook.sheetnames, ["Register"])

    def test_34_phase7_has_all_required_checklist_items(self):
        self.assertEqual(len(default_checklist()), 37)

    def test_35_summary_exposes_delivery_recipient_and_conditions(self):
        config, _artifact, _acceptance_id = self.decided("accepted_with_conditions", "Confirm change window")
        view = summary(config, self.root)["active"]
        self.assertEqual(view["delivery"]["recipient_company"], "Northstar")
        self.assertEqual(view["outstanding_conditions"], 1)


if __name__ == "__main__":
    unittest.main()
