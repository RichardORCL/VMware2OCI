#!/usr/bin/env python3
"""Build the automation-ready OCVS SDD template without touching the source."""

from __future__ import annotations

import argparse
import copy
import hashlib
import re
import shutil
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from lxml import etree


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W15_NS = "http://schemas.microsoft.com/office/word/2012/wordml"
NS = {"w": W_NS, "w15": W15_NS}
ZERO_WIDTH = "\u200b"
DEFAULT_CONTROL_TEXT = {
    "document_version": "0.1",
    "cover_document_version": "0.1",
    "draft_watermark": "DRAFT — NOT FOR CUSTOMER DELIVERY",
}

SOURCE_SHA256 = "209dba06013d74199d07b1c0cb89aba0667e94ec56df1c711d4deaf60bea6817"

SCALAR_TAGS = {
    "customer_name",
    "customer_legal_name",
    "project_name",
    "document_version",
    "generation_date",
    "copyright_year",
    "document_author",
    "document_author_email",
    "document_version_comment",
    "customer_business_context",
    "business_drivers",
    "solution_scope_summary",
    "business_requirements",
    "technical_requirements",
    "success_criteria",
    "rvtools_file_name",
    "assessment_name",
    "assessment_date",
    "selected_vm_count",
    "powered_on_vm_count",
    "powered_off_vm_count",
    "selected_vcpu",
    "selected_ram_gb",
    "selected_storage_tb",
    "operating_system_summary",
    "source_datacenters",
    "source_vcenters",
    "business_scenario",
    "ocvs_topology",
    "sddc_count",
    "target_cluster_count",
    "total_ocvs_nodes",
    "selected_shape",
    "sizing_driver",
    "workload_nodes",
    "spare_nodes",
    "storage_architecture",
    "workload_capacity_storage_tb",
    "vcpu_per_ocpu",
    "cpu_headroom_pct",
    "ram_headroom_pct",
    "storage_headroom_pct",
    "dense_vsan_usable_pct",
    "standard_storage_vpu",
    "oci_region",
    "tenancy_name_or_ocid",
    "compartment_name_or_ocid",
    "vcn_cidr",
    "sddc_cidr",
    "workload_cidrs",
    "connectivity_type",
    "security_requirements",
    "compliance_requirements",
    "identity_federation",
    "logging_siem",
    "encryption_key_management",
    "hcx_required",
    "migration_method",
    "migration_window",
    "downtime_tolerance",
    "validation_and_rollback",
    "ha_requirements",
    "dr_requirements",
    "backup_architecture",
    "operating_model",
    "monitoring_model",
    "implementation_provider",
    "transition_plan",
    "currency_code",
    "iaas_discount_pct",
    "commitment_term",
    "monthly_cost",
    "annual_cost",
    "pricing_availability",
    "draft_watermark",
}

REPEAT_TAGS = {
    "document_change_log",
    "document_reviewers",
    "document_approvers",
    "project_team",
    "environments",
    "source_cluster_rows",
    "target_ocvs_cluster_rows",
    "bom_rows",
    "raci_rows",
    "risk_rows",
    "assumption_rows",
    "customer_obligations",
    "implementation_scope",
    "transition_milestones",
    "network_segments",
}

CONDITIONAL_TAGS = {
    "section_single_cluster",
    "section_multi_cluster",
    "section_vsan",
    "section_block_volume",
    "section_denseio",
    "section_standard_optimized",
    "section_fastconnect",
    "section_ipsec_vpn",
    "section_hcx",
    "section_ha",
    "section_dr",
    "section_backup",
    "section_security_compliance",
    "section_customer_managed",
    "section_partner_managed",
    "section_oracle_managed",
    "section_oracle_lift",
    "section_pricing",
    "section_raci",
    "section_risks",
    "section_transition",
}

CHILD_TAGS = {
    "cover_customer_name",
    "cover_document_version",
    "cover_generation_date",
    "copyright_year_cover",
    "version_history_version",
    "version_history_author",
    "version_history_date",
    "version_history_comment",
    "reviewer_name",
    "reviewer_email",
    "reviewer_role",
    "reviewer_company",
    "approver_name",
    "approver_email",
    "approver_role",
    "approver_company",
    "project_team_name",
    "project_team_email",
    "project_team_role",
    "project_team_company",
    "environment_name",
    "environment_scope",
    "environment_location",
    "environment_assessment_coverage",
    "source_cluster_name",
    "source_cluster_vm_count",
    "source_cluster_vcpu",
    "source_cluster_ram_gb",
    "source_cluster_storage_tb",
    "source_cluster_powered_on",
    "source_cluster_powered_off",
    "target_cluster_name",
    "target_cluster_role",
    "assigned_source_clusters",
    "target_cluster_vm_count",
    "target_cluster_shape",
    "target_cluster_workload_nodes",
    "target_cluster_spare_nodes",
    "target_cluster_total_nodes",
    "target_cluster_sizing_driver",
    "target_cluster_storage_tb",
    "target_cluster_storage_architecture",
    "target_cluster_monthly_cost",
    "bom_part_number",
    "bom_product_name",
    "bom_metric",
    "bom_quantity",
    "bom_commitment_term",
    "bom_storage_component",
    "bom_pricing_availability",
    "raci_activity",
    "raci_responsible",
    "raci_accountable",
    "raci_consulted",
    "raci_informed",
    "risk_id",
    "risk_description",
    "risk_impact",
    "risk_mitigation",
    "risk_owner",
    "assumption_statement",
    "assumption_owner",
    "assumption_status",
    "obligation_statement",
    "obligation_owner",
    "obligation_due_date",
    "scope_activity",
    "scope_deliverable",
    "scope_owner",
    "transition_milestone",
    "transition_owner",
    "transition_date",
    "transition_acceptance",
    "network_segment_name",
    "network_segment_cidr",
    "network_segment_purpose",
    "target_region_summary",
    "single_cluster_shape",
    "single_cluster_workload_nodes",
    "single_cluster_spare_nodes",
    "single_cluster_sizing_driver",
    "single_cluster_storage_architecture",
    "dense_shape",
    "dense_total_nodes",
    "dense_storage_requirement_tb",
    "dense_storage_headroom_pct",
    "dense_vsan_usable_pct_detail",
    "block_shape",
    "block_total_nodes",
    "block_storage_requirement_tb",
    "block_storage_headroom_pct",
    "block_standard_storage_vpu_detail",
    "target_sizing_summary_narrative",
}

ALL_TAGS = SCALAR_TAGS | REPEAT_TAGS | CONDITIONAL_TAGS | CHILD_TAGS


def marker(tag: str) -> str:
    return f"[[SDT:{tag}]]"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_paragraph_text(paragraph, text: str, *, bold: bool = False) -> None:
    clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    run.bold = bold


def set_cell_text(cell, text: str, *, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def remove_row(table, row_index: int) -> None:
    table._tbl.remove(table.rows[row_index]._tr)


def remove_table(table) -> None:
    parent = table._tbl.getparent()
    parent.remove(table._tbl)


def find_paragraph(doc: Document, text: str):
    for p in doc.paragraphs:
        if p.text.strip() == text:
            return p
    raise ValueError(f"Paragraph not found: {text}")


def xml_paragraph(text: str) -> OxmlElement:
    p = OxmlElement("w:p")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    p.append(r)
    return p


def add_section_markers(doc: Document, start_heading: str, end_heading: str, tag: str) -> None:
    start = find_paragraph(doc, start_heading)
    end = find_paragraph(doc, end_heading)
    start._p.addprevious(xml_paragraph(f"[[SECTION:{tag}:START]]"))
    end._p.addprevious(xml_paragraph(f"[[SECTION:{tag}:END]]"))


def insert_table_after(anchor, headers: list[str], tags: list[str], repeat_tag: str | None = None):
    doc = anchor.part.document
    table = doc.add_table(rows=2, cols=len(headers))
    table.style = "Table"
    for idx, value in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], value, bold=True)
    for idx, tag in enumerate(tags):
        set_cell_text(table.rows[1].cells[idx], marker(tag))
    anchor._p.addnext(table._tbl)
    if repeat_tag:
        table._tbl.set(qn("w:description"), repeat_tag)
    return table


def insert_key_value_table_after(anchor, rows: list[tuple[str, str]], description: str | None = None):
    doc = anchor.part.document
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table"
    for row, (label, tag) in zip(table.rows, rows):
        set_cell_text(row.cells[0], label, bold=True)
        set_cell_text(row.cells[1], marker(tag))
    anchor._p.addnext(table._tbl)
    if description:
        table._tbl.set(qn("w:description"), description)
    return table


def insert_heading_and_table_before(doc: Document, before_heading: str, heading_text: str, section_tag: str):
    before = find_paragraph(doc, before_heading)
    end = xml_paragraph(f"[[SECTION:{section_tag}:END]]")
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table"
    set_cell_text(table.cell(0, 0), "Implementation model details are supplied through SDD Configuration.")
    heading = doc.add_paragraph(heading_text, style="Heading 2")
    start = xml_paragraph(f"[[SECTION:{section_tag}:START]]")
    before._p.addprevious(end)
    end.addprevious(table._tbl)
    table._tbl.addprevious(heading._p)
    heading._p.addprevious(start)


def add_draft_header(doc: Document) -> None:
    header = doc.sections[0].header
    p = header.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(marker("draft_watermark"))
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(165, 165, 165)
    r.bold = True


def clean_high_level_document(source: Path, working: Path) -> None:
    doc = Document(source)
    # Capture the source tables before inserting any new tables.  python-docx
    # exposes a live, document-order collection, so positional lookups after an
    # insertion could otherwise target a newly-created table.
    original_tables = list(doc.tables)
    doc.core_properties.title = "Oracle Cloud VMware Solution Definition Document"
    doc.core_properties.subject = "Move to OCVS customer solution definition"
    doc.core_properties.comments = "Automation-ready template"

    # Keep the approved confidentiality footer wording while making its year
    # automation-ready.  The cover uses a separate alias control so scalar tags
    # remain unique.
    for paragraph in doc.sections[0].footer.paragraphs:
        if "Copyright" in paragraph.text:
            clear_paragraph(paragraph)
            paragraph.add_run("Copyright © ")
            paragraph.add_run(marker("copyright_year"))
            paragraph.add_run(", Oracle and/or its affiliates")

    # Cover page text boxes are finalized at XML level; document controls are added here.
    table0 = original_tables[0]
    while len(table0.rows) > 2:
        remove_row(table0, len(table0.rows) - 1)
    for index, tag in enumerate(
        ["version_history_version", "version_history_author", "version_history_date", "version_history_comment"]
    ):
        set_cell_text(table0.rows[1].cells[index], marker(tag))
    table0._tbl.set(qn("w:description"), "document_change_log")

    table1 = original_tables[1]
    while len(table1.rows) > 2:
        remove_row(table1, len(table1.rows) - 1)
    for index, tag in enumerate(["reviewer_name", "reviewer_email", "reviewer_role", "reviewer_company"]):
        set_cell_text(table1.rows[1].cells[index], marker(tag))
    table1._tbl.set(qn("w:description"), "document_reviewers")

    team_heading = find_paragraph(doc, "Team")
    approvers = insert_table_after(
        team_heading,
        ["Approver", "E-Mail", "Role", "Company"],
        ["approver_name", "approver_email", "approver_role", "approver_company"],
        "document_approvers",
    )
    project_team = doc.add_table(rows=2, cols=4)
    project_team.style = "Table"
    for i, value in enumerate(["Project Team Member", "E-Mail", "Role", "Company"]):
        set_cell_text(project_team.rows[0].cells[i], value, bold=True)
    for i, tag in enumerate(["project_team_name", "project_team_email", "project_team_role", "project_team_company"]):
        set_cell_text(project_team.rows[1].cells[i], marker(tag))
    approvers._tbl.addnext(project_team._tbl)
    project_team._tbl.set(qn("w:description"), "project_team")

    # Business narrative: replace samples with content controls rather than visible placeholders.
    set_paragraph_text(doc.paragraphs[12], marker("customer_business_context"))
    set_paragraph_text(doc.paragraphs[13], "")
    set_paragraph_text(doc.paragraphs[14], "")
    set_paragraph_text(doc.paragraphs[18], marker("business_drivers"))
    set_paragraph_text(doc.paragraphs[23], marker("solution_scope_summary"))

    document_purpose = find_paragraph(doc, "Document Purpose")
    metadata = insert_key_value_table_after(
        document_purpose,
        [
            ("Customer name", "customer_name"),
            ("Legal customer name", "customer_legal_name"),
            ("Project name", "project_name"),
            ("Document version", "document_version"),
            ("Generation date", "generation_date"),
            ("Document author", "document_author"),
            ("Author e-mail", "document_author_email"),
            ("Version comment", "document_version_comment"),
            ("Assessment name", "assessment_name"),
            ("Assessment date", "assessment_date"),
            ("RVTools source", "rvtools_file_name"),
            ("Source datacenters", "source_datacenters"),
            ("Source vCenters", "source_vcenters"),
            ("Business scenario", "business_scenario"),
        ],
        "document_metadata",
    )

    # Requirements controls.
    requirements_heading = find_paragraph(doc, "Non-Functional Requirements")
    requirements = insert_key_value_table_after(
        requirements_heading,
        [
            ("Business requirements", "business_requirements"),
            ("Technical requirements", "technical_requirements"),
            ("Success criteria", "success_criteria"),
            ("High availability requirements", "ha_requirements"),
            ("Disaster recovery requirements", "dr_requirements"),
            ("Security requirements", "security_requirements"),
            ("Compliance requirements", "compliance_requirements"),
        ],
        "requirements_summary",
    )

    # Environment rows.
    table3 = original_tables[3]
    while len(table3.rows) > 2:
        remove_row(table3, len(table3.rows) - 1)
    for i, value in enumerate(["Environment", "Scope", "Location", "Assessment coverage"]):
        set_cell_text(table3.rows[0].cells[i], value, bold=True)
    for i, tag in enumerate(["environment_name", "environment_scope", "environment_location", "environment_assessment_coverage"]):
        set_cell_text(table3.rows[1].cells[i], marker(tag))
    table3._tbl.set(qn("w:description"), "environments")

    # Replace current-state sample tables.
    current_inventory_heading = find_paragraph(doc, "Current VMware Inventory On-premises")
    source_table = insert_table_after(
        current_inventory_heading,
        ["Source cluster", "VMs", "vCPU", "RAM GB", "Storage TB", "Powered on", "Powered off"],
        [
            "source_cluster_name",
            "source_cluster_vm_count",
            "source_cluster_vcpu",
            "source_cluster_ram_gb",
            "source_cluster_storage_tb",
            "source_cluster_powered_on",
            "source_cluster_powered_off",
        ],
        "source_cluster_rows",
    )
    total_heading = find_paragraph(doc, "Total VMware Resources For All Locations")
    total_table = insert_table_after(
        total_heading,
        ["Selected VMs", "Powered on", "Powered off", "vCPU", "RAM GB", "Storage TB", "Operating systems"],
        [
            "selected_vm_count",
            "powered_on_vm_count",
            "powered_off_vm_count",
            "selected_vcpu",
            "selected_ram_gb",
            "selected_storage_tb",
            "operating_system_summary",
        ],
    )
    utilization_heading = find_paragraph(doc, "Resource Utilization")
    policy_table = insert_table_after(
        utilization_heading,
        ["vCPU per OCPU", "CPU headroom", "RAM headroom", "Storage headroom", "vSAN usable", "Block Volume VPU/GB"],
        [
            "vcpu_per_ocpu",
            "cpu_headroom_pct",
            "ram_headroom_pct",
            "storage_headroom_pct",
            "dense_vsan_usable_pct",
            "standard_storage_vpu",
        ],
    )
    for old in [original_tables[4], original_tables[5], original_tables[6]]:
        remove_table(old)

    # Network input tables replace sample CIDRs.
    network_heading = find_paragraph(doc, "OCVS Specific Networking Configuration Within OCI")
    network_table = insert_table_after(
        network_heading,
        ["Network segment", "CIDR", "Purpose"],
        ["network_segment_name", "network_segment_cidr", "network_segment_purpose"],
        "network_segments",
    )
    vcn_heading = find_paragraph(doc, "OCI Virtual Cloud Network (VCN) for OCVS")
    vcn_table = insert_key_value_table_after(
        vcn_heading,
        [("VCN CIDR", "vcn_cidr"), ("SDDC CIDR", "sddc_cidr"), ("Workload CIDRs", "workload_cidrs")],
        "network_cidr_summary",
    )
    remove_table(original_tables[8])
    remove_table(original_tables[9])

    # Configuration controls in relevant sections.
    security_heading = find_paragraph(doc, "Security")
    security_table = insert_key_value_table_after(
        security_heading,
        [
            ("Identity federation", "identity_federation"),
            ("Logging and SIEM", "logging_siem"),
            ("Encryption and key management", "encryption_key_management"),
        ],
        "security_configuration",
    )
    networking_heading = find_paragraph(doc, "Networking")
    connectivity_table = insert_key_value_table_after(
        networking_heading,
        [("Connectivity type", "connectivity_type"), ("Target OCI region", "target_region_summary")],
        "connectivity_summary",
    )
    migration_heading = find_paragraph(doc, "OCVS Migration")
    migration_table = insert_key_value_table_after(
        migration_heading,
        [
            ("HCX required", "hcx_required"),
            ("Migration method", "migration_method"),
            ("Migration window", "migration_window"),
            ("Downtime tolerance", "downtime_tolerance"),
            ("Validation and rollback", "validation_and_rollback"),
        ],
        "migration_configuration",
    )
    operations_heading = find_paragraph(doc, "Operating Model, Monitoring, and Management")
    operations_table = insert_key_value_table_after(
        operations_heading,
        [
            ("Operating model", "operating_model"),
            ("Monitoring model", "monitoring_model"),
            ("Backup architecture", "backup_architecture"),
        ],
        "operations_configuration",
    )

    # Replace target sizing and BOM with dynamic structures.
    target_heading = find_paragraph(doc, "Target OCVS Architecture Sizing")
    summary_table = insert_key_value_table_after(
        target_heading,
        [
            ("Target OCI region", "oci_region"),
            ("Tenancy", "tenancy_name_or_ocid"),
            ("Compartment", "compartment_name_or_ocid"),
            ("Topology", "ocvs_topology"),
            ("SDDCs", "sddc_count"),
            ("Target clusters", "target_cluster_count"),
            ("Total OCVS nodes", "total_ocvs_nodes"),
            ("Selected shape", "selected_shape"),
            ("Sizing driver", "sizing_driver"),
            ("Workload nodes", "workload_nodes"),
            ("Spare nodes", "spare_nodes"),
            ("Storage architecture", "storage_architecture"),
            ("Workload storage requirement TB", "workload_capacity_storage_tb"),
        ],
        "target_ocvs_summary",
    )
    target_table = doc.add_table(rows=2, cols=8)
    target_table.style = "Table"
    for i, value in enumerate(["Cluster", "Role", "Assigned sources", "VMs", "Shape", "Nodes", "Storage and architecture", "Monthly cost"]):
        set_cell_text(target_table.rows[0].cells[i], value, bold=True)
    target_cell_tags = [
        ["target_cluster_name"],
        ["target_cluster_role"],
        ["assigned_source_clusters"],
        ["target_cluster_vm_count"],
        ["target_cluster_shape"],
        ["target_cluster_workload_nodes", "target_cluster_spare_nodes", "target_cluster_total_nodes", "target_cluster_sizing_driver"],
        ["target_cluster_storage_tb", "target_cluster_storage_architecture"],
        ["target_cluster_monthly_cost"],
    ]
    for cell, tags in zip(target_table.rows[1].cells, target_cell_tags):
        cell.text = ""
        for index, tag in enumerate(tags):
            p = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
            p.add_run(marker(tag))
    summary_table._tbl.addnext(target_table._tbl)
    target_table._tbl.set(qn("w:description"), "target_ocvs_cluster_rows")

    # Replace sample outcome assertions (fixed node counts, ratios and failure
    # statements) with a rich-text control for the future approved narrative.
    body_paragraphs = list(doc.paragraphs)
    summary_index = next(
        i for i, p in enumerate(body_paragraphs)
        if p.text.strip() == "Summary Of Target Architecture Sizing"
    )
    boq_index = next(
        i for i, p in enumerate(body_paragraphs[summary_index + 1 :], summary_index + 1)
        if p.text.strip() == "BoQ Notes:"
    )
    between = body_paragraphs[summary_index + 1 : boq_index]
    if between:
        set_paragraph_text(between[0], marker("target_sizing_summary_narrative"))
        between[0].style = doc.styles["Normal"]
        for paragraph in between[1:]:
            set_paragraph_text(paragraph, "")
            paragraph.style = doc.styles["Normal"]
        # The source paragraphs carry direct list numbering in addition to
        # their paragraph style.  Resetting only the style leaves visible
        # orphan bullets after their sample text is removed.
        for paragraph in between:
            ppr = paragraph._p.get_or_add_pPr()
            numpr = ppr.find(qn("w:numPr"))
            if numpr is not None:
                ppr.remove(numpr)

    # Single cluster module.
    single_start = xml_paragraph("[[SECTION:section_single_cluster:START]]")
    single_heading = doc.add_paragraph("Consolidated single-cluster sizing", style="Heading 3")
    single_table = doc.add_table(rows=1, cols=5)
    single_table.style = "Table"
    for i, tag in enumerate(["single_cluster_shape", "single_cluster_workload_nodes", "single_cluster_spare_nodes", "single_cluster_sizing_driver", "single_cluster_storage_architecture"]):
        set_cell_text(single_table.rows[0].cells[i], marker(tag))
    single_end = xml_paragraph("[[SECTION:section_single_cluster:END]]")
    target_table._tbl.addnext(single_end)
    single_end.addprevious(single_table._tbl)
    single_table._tbl.addprevious(single_heading._p)
    single_heading._p.addprevious(single_start)

    # Separate storage architecture modules.
    vsan_start = xml_paragraph("[[SECTION:section_vsan:START]]")
    dense_start = xml_paragraph("[[SECTION:section_denseio:START]]")
    vsan_heading = doc.add_paragraph("DenseIO and vSAN sizing", style="Heading 3")
    vsan_table = doc.add_table(rows=2, cols=5)
    vsan_table.style = "Table"
    for i, value in enumerate(["Shape", "Nodes", "Storage requirement TB", "Storage headroom", "Usable vSAN"]):
        set_cell_text(vsan_table.rows[0].cells[i], value, bold=True)
    for i, tag in enumerate(["dense_shape", "dense_total_nodes", "dense_storage_requirement_tb", "dense_storage_headroom_pct", "dense_vsan_usable_pct_detail"]):
        set_cell_text(vsan_table.rows[1].cells[i], marker(tag))
    dense_end = xml_paragraph("[[SECTION:section_denseio:END]]")
    vsan_end = xml_paragraph("[[SECTION:section_vsan:END]]")
    single_end.addnext(vsan_end)
    vsan_end.addprevious(dense_end)
    dense_end.addprevious(vsan_table._tbl)
    vsan_table._tbl.addprevious(vsan_heading._p)
    vsan_heading._p.addprevious(dense_start)
    dense_start.addprevious(vsan_start)

    block_start = xml_paragraph("[[SECTION:section_block_volume:START]]")
    standard_start = xml_paragraph("[[SECTION:section_standard_optimized:START]]")
    block_heading = doc.add_paragraph("Standard or Optimized and OCI Block Volume sizing", style="Heading 3")
    block_table = doc.add_table(rows=2, cols=5)
    block_table.style = "Table"
    for i, value in enumerate(["Shape", "Nodes", "Storage requirement TB", "Storage headroom", "VPU per GB"]):
        set_cell_text(block_table.rows[0].cells[i], value, bold=True)
    for i, tag in enumerate(["block_shape", "block_total_nodes", "block_storage_requirement_tb", "block_storage_headroom_pct", "block_standard_storage_vpu_detail"]):
        set_cell_text(block_table.rows[1].cells[i], marker(tag))
    standard_end = xml_paragraph("[[SECTION:section_standard_optimized:END]]")
    block_end = xml_paragraph("[[SECTION:section_block_volume:END]]")
    vsan_end.addnext(block_end)
    block_end.addprevious(standard_end)
    standard_end.addprevious(block_table._tbl)
    block_table._tbl.addprevious(block_heading._p)
    block_heading._p.addprevious(standard_start)
    standard_start.addprevious(block_start)

    # Multi-cluster marker encloses the repeatable target table only.
    target_table._tbl.addprevious(xml_paragraph("[[SECTION:section_multi_cluster:START]]"))
    target_table._tbl.addnext(xml_paragraph("[[SECTION:section_multi_cluster:END]]"))

    # BOM and pricing.
    sizing_heading = find_paragraph(doc, "Sizing and Bill of Materials")
    bom_table = doc.add_table(rows=2, cols=7)
    bom_table.style = "Table"
    for i, value in enumerate(["Part number", "Product", "Metric", "Quantity", "Commitment", "Storage", "Pricing"]):
        set_cell_text(bom_table.rows[0].cells[i], value, bold=True)
    for i, tag in enumerate(["bom_part_number", "bom_product_name", "bom_metric", "bom_quantity", "bom_commitment_term", "bom_storage_component", "bom_pricing_availability"]):
        set_cell_text(bom_table.rows[1].cells[i], marker(tag))
    sizing_heading._p.addnext(bom_table._tbl)
    bom_table._tbl.set(qn("w:description"), "bom_rows")

    pricing_start = xml_paragraph("[[SECTION:section_pricing:START]]")
    pricing_heading = doc.add_paragraph("Commercial Summary", style="Heading 3")
    pricing_table = doc.add_table(rows=6, cols=2)
    pricing_table.style = "Table"
    for row, (label, tag) in zip(
        pricing_table.rows,
        [
            ("Currency", "currency_code"),
            ("IaaS discount", "iaas_discount_pct"),
            ("Commitment term", "commitment_term"),
            ("Monthly cost", "monthly_cost"),
            ("Annual cost", "annual_cost"),
            ("Pricing availability", "pricing_availability"),
        ],
    ):
        set_cell_text(row.cells[0], label, bold=True)
        set_cell_text(row.cells[1], marker(tag))
    pricing_end = xml_paragraph("[[SECTION:section_pricing:END]]")
    bom_table._tbl.addnext(pricing_end)
    pricing_end.addprevious(pricing_table._tbl)
    pricing_table._tbl.addprevious(pricing_heading._p)
    pricing_heading._p.addprevious(pricing_start)

    for old in [original_tables[10], original_tables[11], original_tables[12]]:
        remove_table(old)

    # Implementation model modules before the Oracle-specific section.
    implementation_heading = find_paragraph(doc, "Project Implementation (Only for Oracle Implementations!)")
    general_heading = doc.add_paragraph("Implementation Approach", style="Heading 1")
    implementation_heading._p.addprevious(general_heading._p)
    insert_heading_and_table_before(doc, "Project Implementation (Only for Oracle Implementations!)", "Customer-managed implementation", "section_customer_managed")
    insert_heading_and_table_before(doc, "Project Implementation (Only for Oracle Implementations!)", "Partner-managed implementation", "section_partner_managed")
    insert_heading_and_table_before(doc, "Project Implementation (Only for Oracle Implementations!)", "Oracle-managed implementation", "section_oracle_managed")

    # Governance repeatable structures replace sample rows.
    raci_heading = find_paragraph(doc, "Implementation RACI")
    raci_table = insert_table_after(
        raci_heading,
        ["Activity", "Responsible", "Accountable", "Consulted", "Informed"],
        ["raci_activity", "raci_responsible", "raci_accountable", "raci_consulted", "raci_informed"],
        "raci_rows",
    )
    assumptions_heading = find_paragraph(doc, "Assumptions")
    assumptions_table = insert_table_after(
        assumptions_heading,
        ["Assumption", "Owner", "Status"],
        ["assumption_statement", "assumption_owner", "assumption_status"],
        "assumption_rows",
    )
    obligations_heading = find_paragraph(doc, "Obligations")
    obligations_table = insert_table_after(
        obligations_heading,
        ["Customer obligation", "Owner", "Due date"],
        ["obligation_statement", "obligation_owner", "obligation_due_date"],
        "customer_obligations",
    )
    risks_heading = find_paragraph(doc, "Risks")
    risks_table = insert_table_after(
        risks_heading,
        ["ID", "Risk", "Impact", "Mitigation", "Owner"],
        ["risk_id", "risk_description", "risk_impact", "risk_mitigation", "risk_owner"],
        "risk_rows",
    )
    workplan_heading = find_paragraph(doc, "Workplan")
    scope_table = insert_table_after(
        workplan_heading,
        ["Activity", "Deliverable", "Owner"],
        ["scope_activity", "scope_deliverable", "scope_owner"],
        "implementation_scope",
    )
    transition_heading = find_paragraph(doc, "Transition Plan")
    transition_table = insert_table_after(
        transition_heading,
        ["Milestone", "Owner", "Date", "Acceptance"],
        ["transition_milestone", "transition_owner", "transition_date", "transition_acceptance"],
        "transition_milestones",
    )
    remove_table(original_tables[13])
    remove_table(original_tables[14])

    # Specialist validation controls in the implementation area.
    provider_table = insert_key_value_table_after(
        general_heading,
        [("Implementation provider", "implementation_provider"), ("Transition plan summary", "transition_plan")],
        "implementation_configuration",
    )

    # Conditional markers around existing modules. Inner sections first at XML patch time.
    add_section_markers(doc, "OCVS Resilience and Recovery", "OCVS High Availability", "section_dr")
    add_section_markers(doc, "OCVS Backup", "Security", "section_backup")
    add_section_markers(doc, "High Availability and Disaster Recovery", "Security", "section_ha")
    add_section_markers(doc, "Security", "Networking", "section_security_compliance")
    add_section_markers(doc, "IPSec VPN", "Fast Connect", "section_ipsec_vpn")
    add_section_markers(doc, "Fast Connect", "OCVS Specific Networking Configuration Within OCI", "section_fastconnect")
    add_section_markers(doc, "OCVS Migration", "Sizing and Bill of Materials", "section_hcx")
    add_section_markers(doc, "Implementation RACI", "Assumptions", "section_raci")
    add_section_markers(doc, "Risks", "Transition Plan", "section_risks")
    add_section_markers(doc, "Transition Plan", "Annex", "section_transition")
    add_section_markers(doc, "Project Implementation (Only for Oracle Implementations!)", "Annex", "section_oracle_lift")

    add_draft_header(doc)
    doc.save(working)


def element_text(element) -> str:
    return "".join(element.xpath(".//w:t/text()", namespaces=NS))


def create_sdt_pr(tag: str, alias: str | None = None, *, repeating: bool = False, item: bool = False):
    pr = etree.Element(f"{{{W_NS}}}sdtPr")
    alias_el = etree.SubElement(pr, f"{{{W_NS}}}alias")
    alias_el.set(f"{{{W_NS}}}val", alias or tag.replace("_", " ").title())
    tag_el = etree.SubElement(pr, f"{{{W_NS}}}tag")
    tag_el.set(f"{{{W_NS}}}val", tag)
    ident = etree.SubElement(pr, f"{{{W_NS}}}id")
    ident.set(f"{{{W_NS}}}val", str(abs(hash((tag, alias))) % 2000000000 + 1))
    if repeating:
        etree.SubElement(pr, f"{{{W15_NS}}}repeatingSection")
    elif item:
        etree.SubElement(pr, f"{{{W15_NS}}}repeatingSectionItem")
    else:
        etree.SubElement(pr, f"{{{W_NS}}}text")
    return pr


def wrap_text_markers(root) -> set[str]:
    found: set[str] = set()
    for text_node in list(root.xpath(".//w:t", namespaces=NS)):
        raw = text_node.text or ""
        match = re.fullmatch(r"\[\[SDT:([a-z0-9_]+)\]\]", raw)
        if not match:
            continue
        tag = match.group(1)
        found.add(tag)
        run = text_node.getparent()
        parent = run.getparent()
        if parent is None:
            continue
        index = parent.index(run)
        new_run = copy.deepcopy(run)
        new_text = new_run.find(f".//{{{W_NS}}}t")
        if new_text is not None:
            new_text.text = DEFAULT_CONTROL_TEXT.get(tag, ZERO_WIDTH)
            if tag in DEFAULT_CONTROL_TEXT:
                new_text.set(f"{{http://www.w3.org/XML/1998/namespace}}space", "preserve")
        sdt = etree.Element(f"{{{W_NS}}}sdt")
        sdt.append(create_sdt_pr(tag))
        content = etree.SubElement(sdt, f"{{{W_NS}}}sdtContent")
        content.append(new_run)
        parent.remove(run)
        parent.insert(index, sdt)
    return found


def wrap_repeat_row(table, repeat_tag: str) -> None:
    rows = table.xpath("./w:tr", namespaces=NS)
    if len(rows) < 2:
        raise ValueError(f"Repeat table {repeat_tag} has no item row")
    row = rows[-1]
    parent = row.getparent()
    idx = parent.index(row)
    outer = etree.Element(f"{{{W_NS}}}sdt")
    outer.append(create_sdt_pr(repeat_tag, repeating=True))
    outer_content = etree.SubElement(outer, f"{{{W_NS}}}sdtContent")
    item = etree.SubElement(outer_content, f"{{{W_NS}}}sdt")
    item.append(create_sdt_pr(f"{repeat_tag}_item", item=True))
    item_content = etree.SubElement(item, f"{{{W_NS}}}sdtContent")
    parent.remove(row)
    item_content.append(row)
    parent.insert(idx, outer)


def find_marker_paragraph(root, marker_text: str):
    for p in root.xpath(".//w:p", namespaces=NS):
        if element_text(p).strip() == marker_text:
            return p
    return None


def wrap_section(root, tag: str) -> None:
    start_text = f"[[SECTION:{tag}:START]]"
    end_text = f"[[SECTION:{tag}:END]]"
    start = find_marker_paragraph(root, start_text)
    end = find_marker_paragraph(root, end_text)
    if start is None or end is None:
        raise ValueError(f"Conditional section markers not found: {tag}")
    parent = start.getparent()
    if end.getparent() is not parent:
        raise ValueError(f"Conditional section markers have different parents: {tag}")
    children = list(parent)
    start_idx = children.index(start)
    end_idx = children.index(end)
    if end_idx <= start_idx:
        raise ValueError(f"Invalid conditional section range: {tag}")
    sdt = etree.Element(f"{{{W_NS}}}sdt")
    sdt.append(create_sdt_pr(tag))
    content = etree.SubElement(sdt, f"{{{W_NS}}}sdtContent")
    for child in children[start_idx + 1 : end_idx]:
        parent.remove(child)
        content.append(child)
    parent.remove(start)
    parent.remove(end)
    parent.insert(start_idx, sdt)


def hide_conditional_section_by_default(root, tag: str) -> None:
    """Hide a retained conditional module until a future generator enables it."""
    sections = root.xpath(
        f'.//w:sdt[w:sdtPr/w:tag[@w:val="{tag}"]]', namespaces=NS
    )
    if len(sections) != 1:
        raise ValueError(f"Conditional section not found for default hiding: {tag}")
    for run in sections[0].xpath('.//w:r', namespaces=NS):
        rpr = run.find(f'{{{W_NS}}}rPr')
        if rpr is None:
            rpr = etree.Element(f'{{{W_NS}}}rPr')
            run.insert(0, rpr)
        if rpr.find(f'{{{W_NS}}}vanish') is None:
            etree.SubElement(rpr, f'{{{W_NS}}}vanish')
    # Hiding only the runs leaves list numbers, paragraph marks and table rows
    # visible.  Hide paragraph marks and table rows as well so the retained
    # module consumes no default Draft layout space.
    for paragraph in sections[0].xpath('.//w:p', namespaces=NS):
        ppr = paragraph.find(f'{{{W_NS}}}pPr')
        if ppr is None:
            ppr = etree.Element(f'{{{W_NS}}}pPr')
            paragraph.insert(0, ppr)
        mark_rpr = ppr.find(f'{{{W_NS}}}rPr')
        if mark_rpr is None:
            mark_rpr = etree.SubElement(ppr, f'{{{W_NS}}}rPr')
        if mark_rpr.find(f'{{{W_NS}}}vanish') is None:
            etree.SubElement(mark_rpr, f'{{{W_NS}}}vanish')
        numpr = ppr.find(f'{{{W_NS}}}numPr')
        if numpr is not None:
            ppr.remove(numpr)
        style = ppr.find(f'{{{W_NS}}}pStyle')
        if style is None:
            style = etree.SubElement(ppr, f'{{{W_NS}}}pStyle')
        style.set(f'{{{W_NS}}}val', 'Normal')
    for row in sections[0].xpath('.//w:tr', namespaces=NS):
        trpr = row.find(f'{{{W_NS}}}trPr')
        if trpr is None:
            trpr = etree.Element(f'{{{W_NS}}}trPr')
            row.insert(0, trpr)
        if trpr.find(f'{{{W_NS}}}hidden') is None:
            etree.SubElement(trpr, f'{{{W_NS}}}hidden')


def replace_text_nodes(root, replacements: list[tuple[str, str]]) -> None:
    for t in root.xpath(".//w:t", namespaces=NS):
        value = t.text or ""
        for old, new in replacements:
            value = value.replace(old, new)
        t.text = value


def patch_cover_controls(root) -> None:
    """Replace reviewed cover-page sample values with uniquely tagged controls."""
    nodes = root.xpath(".//w:t", namespaces=NS)
    values = [(node.text or "") for node in nodes]

    def replace_first(value: str, replacement: str) -> int:
        for index, current in enumerate(values):
            if current == value:
                nodes[index].text = replacement
                values[index] = replacement
                return index
        raise ValueError(f"Cover value not found: {value}")

    replace_first("Customer Name", marker("cover_customer_name"))

    copyright_index = replace_first("Copyright © 202", "Copyright © ")
    if copyright_index + 1 >= len(nodes) or values[copyright_index + 1] != "4":
        raise ValueError("Reviewed cover copyright year structure changed")
    nodes[copyright_index + 1].text = marker("copyright_year_cover")
    values[copyright_index + 1] = marker("copyright_year_cover")

    date_index = replace_first("6", marker("cover_generation_date"))
    if values[date_index + 1 : date_index + 3] != ["th", " of March 2024"]:
        raise ValueError("Reviewed cover date structure changed")
    nodes[date_index + 1].text = ""
    nodes[date_index + 2].text = ""
    values[date_index + 1] = ""
    values[date_index + 2] = ""

    # The first 1.0 text node after the cover's Version label is the cover value.
    version_label = values.index("Version")
    for index in range(version_label + 1, min(version_label + 8, len(values))):
        if values[index] == "1.0":
            nodes[index].text = marker("cover_document_version")
            values[index] = marker("cover_document_version")
            break
    else:
        raise ValueError("Reviewed cover version value not found")


def patch_package(working: Path, output: Path) -> dict[str, object]:
    with tempfile.TemporaryDirectory(prefix="ocvs-sdd-package-") as tmp:
        tmp_path = Path(tmp)
        with zipfile.ZipFile(working) as archive:
            archive.extractall(tmp_path)

        parser = etree.XMLParser(remove_blank_text=False)
        document_path = tmp_path / "word" / "document.xml"
        document = etree.parse(str(document_path), parser)
        root = document.getroot()
        patch_cover_controls(root)

        # Remove/neutralize engagement samples without creating false customer facts.
        replacements = [
            ("A Company Making Everything's", "the Customer's"),
            ("A Company Making Everything’s", "the Customer's"),
            ("A Company Making Everything\"", "the Customer's"),
            ("A Company Making Everything", "the Customer"),
            ("Oracle OCI Frankfurt cloud region", "the selected OCI region"),
            ("OCI Frankfurt Region", "the selected OCI region"),
            ("Frankfurt, Germany", "the customer location"),
            ("Frankfurt", "the selected OCI region"),
            ("FY23", "the approved programme period"),
            ("EXAMPLE", "the customer programme"),
            ("INSERT REGION", ""),
            ("DC LOCATION", ""),
            ("HARDWARE MODELS", ""),
            ("Name Surname", ""),
            ("example@example.com", ""),
            ("September 01st, 2022", ""),
            ("6th of March 2024", ""),
            ("March 6, 2024", ""),
            (
                "Note: The table is a representation of the sample network layout for the /22 network however based on the SDDC CIDR range selected we can expect an automated network layout created by the SDDC provisioning operation with different network range for VLANS. SDDC workload CIDR will be identified as a part of the low-level design during the implementation.",
                "The diagram illustrates a reference network layout. The final SDDC CIDR and workload network ranges must be validated during the detailed design phase.",
            ),
            (
                "Sample sizing of the OCVS environment is depicted in the table below.",
                "The sizing modules below are populated from the approved OCVS sizing result.",
            ),
        ]
        replace_text_nodes(root, replacements)

        found_tags = wrap_text_markers(root)

        # Convert marked tables into native repeating-section controls.
        repeat_found = set()
        for table in root.xpath(".//w:tbl", namespaces=NS):
            description = table.get(f"{{{W_NS}}}description")
            if description in REPEAT_TAGS:
                del table.attrib[f"{{{W_NS}}}description"]
                wrap_repeat_row(table, description)
                repeat_found.add(description)

        # Inner conditional modules before outer modules.
        section_order = [
            "section_dr",
            "section_backup",
            "section_ipsec_vpn",
            "section_fastconnect",
            "section_raci",
            "section_risks",
            "section_transition",
            "section_denseio",
            "section_standard_optimized",
            "section_single_cluster",
            "section_multi_cluster",
            "section_vsan",
            "section_block_volume",
            "section_pricing",
            "section_customer_managed",
            "section_partner_managed",
            "section_oracle_managed",
            "section_hcx",
            "section_ha",
            "section_security_compliance",
            "section_oracle_lift",
        ]
        for tag in section_order:
            wrap_section(root, tag)

        # Delivery-provider modules are retained for explicit future enablement,
        # but must not appear as customer facts in the default Draft template.
        for hidden_tag in (
            "section_customer_managed",
            "section_partner_managed",
            "section_oracle_managed",
            "section_oracle_lift",
        ):
            hide_conditional_section_by_default(root, hidden_tag)

        document.write(str(document_path), xml_declaration=True, encoding="UTF-8", standalone="yes")

        # Header/footer controls and metadata cleanup.
        other_parts = [
            tmp_path / "word" / "header1.xml",
            tmp_path / "word" / "footer1.xml",
            tmp_path / "word" / "footer2.xml",
        ]
        for part in other_parts:
            if not part.exists():
                continue
            tree = etree.parse(str(part), parser)
            part_root = tree.getroot()
            replace_text_nodes(
                part_root,
                [
                    ("Copyright ©2022", "Copyright ©"),
                    ("Copyright © 2022", "Copyright © "),
                    ("Copyright ©2024", "Copyright ©"),
                    ("Copyright © 2024", "Copyright © "),
                    ("A Company Making Everything", "the Customer"),
                ],
            )
            found_tags |= wrap_text_markers(part_root)
            tree.write(str(part), xml_declaration=True, encoding="UTF-8", standalone="yes")

        # Add update-on-open behavior for TOC/page fields.
        settings_path = tmp_path / "word" / "settings.xml"
        settings = etree.parse(str(settings_path), parser)
        settings_root = settings.getroot()
        for old in settings_root.findall(f"{{{W_NS}}}updateFields"):
            settings_root.remove(old)
        update = etree.SubElement(settings_root, f"{{{W_NS}}}updateFields")
        update.set(f"{{{W_NS}}}val", "true")
        settings.write(str(settings_path), xml_declaration=True, encoding="UTF-8", standalone="yes")

        # The reviewed source uses no separator after the top-level automatic
        # heading number, which renders headings such as ``4Implementation``.
        # Preserve automatic numbering while making the visible result read
        # correctly as ``4 Implementation``.
        numbering_path = tmp_path / "word" / "numbering.xml"
        numbering = etree.parse(str(numbering_path), parser)
        numbering_root = numbering.getroot()
        for level in numbering_root.xpath(
            './/w:abstractNum/w:lvl[@w:ilvl="0"][w:pStyle[@w:val="Heading1"]]',
            namespaces=NS,
        ):
            suffix = level.find(f"{{{W_NS}}}suff")
            if suffix is None:
                suffix = etree.SubElement(level, f"{{{W_NS}}}suff")
            suffix.set(f"{{{W_NS}}}val", "space")
        numbering.write(
            str(numbering_path), xml_declaration=True, encoding="UTF-8", standalone="yes"
        )

        # Ensure repeat namespace is declared through mc:Ignorable if present.
        for path in [document_path]:
            tree = etree.parse(str(path), parser)
            r = tree.getroot()
            ignorable = r.get("{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable", "")
            prefixes = set(ignorable.split())
            prefixes.add("w15")
            r.set("{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable", " ".join(sorted(prefixes)))
            tree.write(str(path), xml_declaration=True, encoding="UTF-8", standalone="yes")

        output.parent.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
            for path in sorted(tmp_path.rglob("*")):
                if path.is_file():
                    archive.write(path, path.relative_to(tmp_path).as_posix())

    return {"scalar_found": sorted(found_tags), "repeat_found": sorted(repeat_found)}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    if sha256(args.source) != SOURCE_SHA256:
        raise SystemExit("Source template checksum differs from the reviewed template")
    with tempfile.TemporaryDirectory(prefix="ocvs-sdd-working-") as tmp:
        working = Path(tmp) / "working.docx"
        clean_high_level_document(args.source, working)
        report = patch_package(working, args.output)
    print(f"Created {args.output}")
    print(f"Scalar/child controls found: {len(report['scalar_found'])}")
    print(f"Repeat structures found: {len(report['repeat_found'])}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
